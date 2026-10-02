"""Classify downloaded Clio documents and extract text from non-PDF files.

PDFs and images go through ingest/pdf.py (images are converted to PDF so the
source viewer can show them and OCR yields line boxes). Everything else is
plain text here. Unknown types return a visible "not processed" marker so a
gap in the record is never silent.
"""

import email
import html
import io
import re
from email import policy

TEXT_PARSER_VERSION = 1  # bump to re-extract non-PDF documents on next sync

IMAGE_EXTS = {"png", "jpg", "jpeg", "tif", "tiff", "gif", "bmp", "webp"}
TEXT_EXTS = {"txt", "text", "md", "csv", "log", "rtf"}
HTML_EXTS = {"html", "htm", "xhtml"}
EMAIL_EXTS = {"eml"}
DOCX_EXTS = {"docx"}


def kind_of(ext: str, ctype: str, head: bytes = b"") -> str:
    """'pdf' | 'image' | 'docx' | 'html' | 'email' | 'text' | 'other'."""
    ext, ctype = (ext or "").lower(), (ctype or "").lower()
    if head[:5] == b"%PDF-" or "pdf" in ctype or ext == "pdf":
        return "pdf"
    if ext in IMAGE_EXTS or ctype.startswith("image/"):
        return "image"
    if ext in DOCX_EXTS or "wordprocessingml" in ctype:
        return "docx"
    if ext in HTML_EXTS or "html" in ctype:
        return "html"
    if ext in EMAIL_EXTS or ctype == "message/rfc822":
        return "email"
    if ext in TEXT_EXTS or ctype.startswith("text/"):
        return "text"
    return "other"


def _decode(data: bytes) -> str:
    for enc in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", "replace")


def _strip_html(s: str) -> str:
    s = re.sub(r"(?is)<(script|style|head)\b.*?</\1>", " ", s)
    s = re.sub(r"(?i)<br\s*/?>|</(p|div|li|tr|h[1-6])>", "\n", s)
    s = re.sub(r"(?s)<[^>]+>", " ", s)
    s = html.unescape(s)
    lines = [" ".join(line.split()) for line in s.splitlines()]
    return "\n".join(line for line in lines if line).strip()


def _docx(data: bytes) -> str:
    import docx

    d = docx.Document(io.BytesIO(data))
    parts = [p.text for p in d.paragraphs if p.text.strip()]
    for table in d.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def _email(data: bytes) -> str:
    msg = email.message_from_bytes(data, policy=policy.default)
    head = "\n".join(f"{h}: {msg[h]}" for h in ("Subject", "From", "To", "Cc", "Date") if msg[h])
    body = msg.get_body(preferencelist=("plain", "html"))
    text = ""
    if body is not None:
        text = body.get_content()
        if body.get_content_type() == "text/html":
            text = _strip_html(text)
    names = [p.get_filename() for p in msg.iter_attachments() if p.get_filename()]
    if names:
        head += f"\nAttachments: {', '.join(names)}"
    return f"{head}\n\n{text}".strip()


def extract_text(kind: str, data: bytes, type_label: str) -> str:
    """Text for a non-PDF, non-image document; marker text for unsupported types."""
    if kind == "docx":
        return _docx(data)
    if kind == "html":
        return _strip_html(_decode(data))
    if kind == "email":
        return _email(data)
    if kind == "text":
        return html.unescape(_decode(data)).strip()
    return f"File type not processed: {type_label}"
